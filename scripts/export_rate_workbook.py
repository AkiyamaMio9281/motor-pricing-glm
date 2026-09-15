"""Writes exports/rate_workbook.xlsx and the Power BI extract in exports/powerbi/.

    .venv/Scripts/python scripts/export_rate_workbook.py
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import pure_premium as pp
import rate_table as rt
from db import dsn

REPO_ROOT = Path(__file__).resolve().parent.parent
EXPORTS = REPO_ROOT / "exports"
WORKBOOK = EXPORTS / "rate_workbook.xlsx"
POWERBI = EXPORTS / "powerbi"
HEADER_FILL = PatternFill("solid", fgColor="E1E0D9")
HEADER_FONT = Font(bold=True, color="0B0B0B")

SEGMENT_COLUMNS = [
    ("region", "Region", None),
    ("driv_age_band", "Driver age band", None),
    ("exposure", "Exposure, policy-years", "#,##0.0"),
    ("reported_claims", "Reported claims", "#,##0"),
    ("priced_claims", "Priced claims", "#,##0"),
    ("incurred_loss", "Incurred loss", "#,##0"),
    ("expected", "GLM manual premium", "#,##0"),
    ("loss_ratio", "Loss ratio, incurred over manual premium", "0.0%"),
    ("actual", "Capped losses times load", "#,##0"),
    ("manual_rate", "GLM manual rate", "#,##0.00"),
    ("experience_rate", "Experience rate", "#,##0.00"),
    ("limited_fluctuation_z", "Limited-fluctuation Z", "0.000"),
    ("buhlmann_z", "Buhlmann Z", "0.000"),
    ("weighted_rate", "Credibility-weighted rate", "#,##0.00"),
    ("normalized_rate", "Normalized rate", "#,##0.00"),
]
RELATIVITY_COLUMNS = [
    ("factor", "Rating factor", None),
    ("level", "Level", None),
    ("exposure_share", "Share of exposure", "0.0%"),
    ("relativity_largest_base", "Relativity, largest-exposure base", "0.000"),
    ("largest_exposure_base", "Largest-exposure base level", None),
    ("relativity_default_base", "Relativity, first-level base", "0.000"),
    ("default_base", "First-level base", None),
]


def write_table(sheet, frame: pd.DataFrame, columns: list[tuple[str, str, str | None]], start_row: int = 1) -> None:
    for j, (_, title, _) in enumerate(columns, start=1):
        cell = sheet.cell(row=start_row, column=j, value=title)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, record in enumerate(frame[[c for c, _, _ in columns]].itertuples(index=False), start=start_row + 1):
        for j, ((_, _, number_format), value) in enumerate(zip(columns, record), start=1):
            cell = sheet.cell(row=i, column=j, value=value.item() if isinstance(value, np.generic) else value)
            if number_format:
                cell.number_format = number_format
    for j, (_, title, _) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(j)].width = max(12, min(30, len(title) + 2))
    sheet.freeze_panes = sheet.cell(row=start_row + 1, column=3)
    sheet.auto_filter.ref = f"A{start_row}:{get_column_letter(len(columns))}{start_row + len(frame)}"


def workbook(result: rt.RateTable, frame_md5: str) -> None:
    b, c = result.base, result.credibility
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    rows = [
        ("Motor TPL rate workbook", None, None),
        ("Source", "freMTPL2, frame md5 " + frame_md5, None),
        ("Pricing model", f"GLM run {pp.PRICING_RUN}: Poisson frequency of priced claims times capped Gamma severity", None),
        ("Base rate, largest-exposure levels, per policy-year", b["largest"], "#,##0.00"),
        ("Reference density for the base rate", b["reference_density"], "#,##0"),
        ("Density relativity exponent", b["density_slope"], "0.0000"),
        ("Large-loss load, claims capped at 34,377", b["large_loss_load"], "0.0000"),
        ("Off-balance factor", b["off_balance"], "0.000000"),
        ("Full credibility standard, claims (5%, 90%, capped severity CV)", c["full_credibility_claims"], "#,##0"),
        ("Capped severity coefficient of variation", c["severity_cv"], "0.000"),
        ("Buhlmann k against the portfolio mean, policy-years", c["portfolio"]["k"], "#,##0"),
        ("Buhlmann between-segment variance against the GLM", c["against_model"]["vhm"], "0.000"),
        ("Normalization factor for credibility-weighted rates", c["normalization_factor"], "0.0000"),
        ("Premium", "base rate x product of relativities x (density / reference density) ^ exponent", None),
        ("Details", "RESULTS.md and docs/credibility.md in the repository", None),
    ]
    for i, (label, value, number_format) in enumerate(rows, start=1):
        summary.cell(row=i, column=1, value=label).font = Font(bold=True)
        cell = summary.cell(row=i, column=2, value=value)
        if number_format:
            cell.number_format = number_format
    summary.column_dimensions["A"].width = 62
    summary.column_dimensions["B"].width = 80

    segments = book.create_sheet("Segments")
    write_table(segments, result.segments.sort_values(["region", "driv_age_band"]), SEGMENT_COLUMNS)
    drilldown = book.create_sheet("Drilldown")
    write_table(drilldown, result.relativities, RELATIVITY_COLUMNS)
    EXPORTS.mkdir(exist_ok=True)
    book.save(WORKBOOK)


def powerbi_extract(conn: psycopg.Connection, result: rt.RateTable) -> dict[str, int]:
    POWERBI.mkdir(parents=True, exist_ok=True)
    frame = result.frame
    queries = {
        "dim_region": "SELECT region_key, region_code FROM dim.region ORDER BY region_key",
        "dim_area": "SELECT area_key, area_code, density_min, density_max FROM dim.area ORDER BY area_key",
        "dim_vehicle": "SELECT vehicle_key, veh_brand, veh_gas, veh_power FROM dim.vehicle ORDER BY vehicle_key",
        "dim_rating_band": "SELECT factor, lower_bound, label FROM mart.rating_band ORDER BY factor, lower_bound",
        "fact_claim": "SELECT claim_key, idpol, claim_amount::float8 AS claim_amount FROM fact.claim ORDER BY claim_key",
        "keys": "SELECT idpol, region_key, area_key, vehicle_key FROM fact.exposure ORDER BY idpol",
    }
    tables = {}
    with conn.cursor() as cur:
        for name, sql in queries.items():
            cur.execute(sql)
            tables[name] = pd.DataFrame(cur.fetchall(), columns=[d.name for d in cur.description])
    keys = tables.pop("keys")
    if not (keys["idpol"].to_numpy() == frame["idpol"].to_numpy()).all():
        raise ValueError("fact.exposure keys do not match the frame policy for policy")
    tables["fact_policy_year"] = pd.DataFrame({
        "idpol": frame["idpol"],
        "region_key": keys["region_key"],
        "area_key": keys["area_key"],
        "vehicle_key": keys["vehicle_key"],
        "driv_age_band": frame["driv_age_band"],
        "veh_age_band": frame["veh_age_band"],
        "bonus_malus_band": frame["bonus_malus_band"],
        "density": frame["density"],
        "exposure": frame["exposure"],
        "reported_claims": frame["claim_nb"],
        "priced_claims": frame["priced_claim_nb"],
        "incurred_loss": frame["incurred_loss"],
        "capped_loss": frame["capped_loss"],
        "glm_pure_premium": frame["amount_per_year"],
        "glm_expected_loss": frame["expected_loss"],
        "in_holdout": frame["risk_group_holdout"],
    })
    tables["rate_relativity"] = result.relativities
    tables["segment_rate"] = result.segments
    rows = {}
    for name, table in tables.items():
        table.to_csv(POWERBI / f"{name}.csv", index=False, lineterminator="\n")
        rows[name] = len(table)
    return rows


def main() -> None:
    with psycopg.connect(dsn()) as conn:
        result = rt.build(conn)
        rows = powerbi_extract(conn, result)
    workbook(result, pp.canonical_md5(result.frame))
    print(f"wrote {WORKBOOK}")
    for name, count in rows.items():
        print(f"wrote {POWERBI / (name + '.csv')}: {count:,} rows")


if __name__ == "__main__":
    main()
