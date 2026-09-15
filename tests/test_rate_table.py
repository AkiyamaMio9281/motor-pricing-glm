from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

import credibility as cr
import export_rate_workbook as ex
import rate_table as rt
from conftest import scalar

REPO_ROOT = Path(__file__).resolve().parent.parent


def test_the_classic_standard_is_the_frequency_only_case():
    assert cr.full_credibility_standard(0.05, 0.90) == pytest.approx(1082.2, abs=0.1)
    assert cr.full_credibility_standard(0.05, 0.90, severity_cv=2.0) == pytest.approx(1082.2 * 5, abs=0.5)
    assert cr.full_credibility_standard(0.025, 0.90) == pytest.approx(4 * cr.full_credibility_standard(0.05, 0.90))


def test_partial_credibility_follows_the_square_root_rule():
    assert list(cr.limited_fluctuation_credibility([0, 270.55, 1082.2, 5000], 1082.2)) == pytest.approx([0, 0.5, 1, 1], abs=1e-4)


def test_buhlmann_straub_gives_no_credibility_to_cells_that_do_not_differ():
    rng = np.random.default_rng(4)
    cell = np.repeat(np.arange(50), 200)
    x = rng.gamma(0.5, 200.0, cell.size)
    weight = rng.uniform(0.2, 1.0, cell.size)
    fit = cr.buhlmann_straub(x, weight, cell)
    assert fit["vhm"] < 0.1 * fit["epv"] / 200


def test_buhlmann_straub_credits_cells_that_do_differ_and_bigger_cells_more():
    rng = np.random.default_rng(5)
    cell = np.repeat(np.arange(50), 400)
    means = rng.gamma(4.0, 25.0, 50)
    x = rng.gamma(2.0, means[cell] / 2.0)
    weight = np.ones(cell.size)
    fit = cr.buhlmann_straub(x, weight, cell)
    assert fit["vhm"] == pytest.approx(np.var(means), rel=0.35)
    z = cr.buhlmann_credibility([10, 400, 4000], fit["k"])
    assert 0 < z[0] < z[1] < z[2] < 1


def test_the_mart_bands_are_the_modelling_bands(staged):
    source = (REPO_ROOT / "R" / "bands.R").read_text(encoding="utf-8")
    for factor, constant in (("driv_age", "DRIVER_AGE_LOWER"), ("veh_age", "VEH_AGE_LOWER"), ("bonus_malus", "BONUS_MALUS_LOWER")):
        lower = [int(v) for v in re.search(rf"^{constant} <- c\((.*?)\)", source, re.M).group(1).split(",")]
        upper = [b - 1 for b in lower[1:]] + [None]
        labels = [f"{lo}+" if hi is None else (str(lo) if hi == lo else f"{lo}-{hi}") for lo, hi in zip(lower, upper)]
        with staged.cursor() as cur:
            cur.execute("SELECT lower_bound, label FROM mart.rating_band WHERE factor = %s ORDER BY lower_bound", (factor,))
            assert cur.fetchall() == list(zip(lower, labels)), f"mart.rating_band differs from {constant} in R/bands.R"


def test_the_segment_views_lose_and_duplicate_nothing(staged):
    assert scalar(staged, "SELECT count(*) FROM mart.policy_segment") == scalar(staged, "SELECT count(*) FROM fact.exposure")
    assert scalar(staged, "SELECT count(*) FROM mart.policy_segment WHERE driv_age_band IS NULL OR veh_age_band IS NULL "
                          "OR bonus_malus_band IS NULL") == 0
    assert scalar(staged, "SELECT sum(priced_claims) FROM mart.experience_by_segment") == scalar(staged, "SELECT count(*) FROM fact.claim")
    assert scalar(staged, "SELECT round(sum(incurred_loss)::numeric, 2) FROM mart.experience_by_segment") == scalar(
        staged, "SELECT sum(claim_amount) FROM fact.claim")


@pytest.fixture(scope="module")
def table(staged):
    if scalar(staged, "SELECT count(*) FROM model.glm_run WHERE run = 'priced_claims'") == 0:
        pytest.skip("no GLM runs; run Rscript R/export_predictions.R")
    return rt.build(staged)


def test_the_rate_table_reprices_every_policy_under_either_base(table):
    for which in ("default", "largest"):
        premium = rt.premium_from_table(table.frame, table.relativities, table.base, which)
        assert np.allclose(premium, table.frame["amount_per_year"], rtol=1e-9), f"the {which}-base table misprices"


def test_the_readable_base_is_the_largest_exposure_level_of_each_factor(table):
    t = table.relativities
    for factor, levels in t.groupby("factor"):
        largest = levels.loc[levels["exposure_share"].idxmax(), "level"]
        assert (levels["largest_exposure_base"] == largest).all()
        assert levels.loc[levels["level"] == largest, "relativity_largest_base"].item() == pytest.approx(1)
    young = t[(t["factor"] == "driv_age_band") & (t["level"] == "18-19")]
    assert young["default_base"].item() == "18-19" and young["exposure_share"].item() < 0.01


def test_the_experience_columns_reconcile_to_the_mart(table, staged):
    s = table.segments
    assert s["reported_claims"].sum() == scalar(staged, "SELECT sum(reported_claims) FROM mart.experience_by_segment")
    assert s["incurred_loss"].sum() == pytest.approx(float(scalar(staged, "SELECT sum(incurred_loss) FROM mart.experience_by_segment")))
    assert s["incurred_loss"].sum() / s["expected"].sum() == pytest.approx(1, abs=1e-9)


def test_normalized_rates_reproduce_the_losses(table):
    s = table.segments
    assert (s["normalized_rate"] * s["exposure"]).sum() == pytest.approx(s["actual"].sum(), rel=1e-12)
    assert table.credibility["against_model"]["vhm"] < 0 and (s["model_complement_z"] == 0).all()


def test_the_credibility_document_states_the_computed_standard(table):
    text = (REPO_ROOT / "docs" / "credibility.md").read_text(encoding="utf-8")
    assert f"**{table.credibility['full_credibility_claims']:,.0f} claims**" in text
    assert f"k = {table.credibility['portfolio']['k']:,.0f} policy-years" in text


def test_the_workbook_has_its_three_sheets(table, tmp_path, monkeypatch):
    monkeypatch.setattr(ex, "EXPORTS", tmp_path)
    monkeypatch.setattr(ex, "WORKBOOK", tmp_path / "rate_workbook.xlsx")
    ex.workbook(table, "0" * 32)
    book = load_workbook(tmp_path / "rate_workbook.xlsx")
    assert book.sheetnames == ["Summary", "Segments", "Drilldown"]
    assert book["Segments"].max_row == len(table.segments) + 1
    assert book["Drilldown"].freeze_panes == "C2"
